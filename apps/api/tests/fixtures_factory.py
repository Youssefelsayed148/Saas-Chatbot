"""Runtime fixture builders for P2-02 parser tests (PDF bytes, DOCX bytes).

PDFs are built as minimal valid PDF 1.4 documents by hand so the suite has
no binary fixtures to maintain; DOCX bytes come from python-docx itself.
"""

import io
import re
import zipfile


def _pdf_escape(line: str) -> str:
    return line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_pdf_pages(page_texts: list[list[str]]) -> bytes:
    """A minimal text PDF with one page per entry; each entry is a list of lines."""
    body = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}

    def emit(object_id: int, data: bytes) -> None:
        offsets[object_id] = len(body)
        body.extend(f"{object_id} 0 obj\n".encode())
        body.extend(data)
        body.extend(b"\nendobj\n")

    # Object numbering: 1 catalog, 2 pages, 3 font, 4.. page/content pairs.
    page_count = len(page_texts)
    font_id = 3
    page_ids = [4 + index * 2 for index in range(page_count)]
    content_ids = [page_id + 1 for page_id in page_ids]

    for index, lines in enumerate(page_texts):
        stream_ops = "BT /F1 12 Tf 72 720 Td\n"
        for line in lines:
            stream_ops += f"({_pdf_escape(line)}) Tj T*\n"
        stream_ops += "ET"
        stream = stream_ops.encode("latin-1", errors="replace")
        emit(
            content_ids[index],
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        )
        emit(
            page_ids[index],
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_ids[index]} 0 R >>"
            ).encode(),
        )
    emit(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    emit(2, f"<< /Type /Pages /Count {page_count} /Kids [{kids}] >>".encode())
    emit(
        font_id,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    )
    xref_offset = len(body)
    max_id = content_ids[-1] if content_ids else 4
    body.extend(f"xref\n0 {max_id + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for object_id in range(1, max_id + 1):
        body.extend(f"{offsets[object_id]:010d} 00000 n \n".encode())
    body.extend(b"trailer\n")
    body.extend(f"<< /Size {max_id + 1} /Root 1 0 R >>\n".encode())
    body.extend(b"startxref\n" + str(xref_offset).encode() + b"\n%%EOF")
    return bytes(body)


def make_scanned_pdf(pages: int = 1) -> bytes:
    """A 'scanned' PDF: pages exist but no text operator produces text."""
    body = bytearray(make_pdf_pages([["", ""] for _ in range(pages)]))
    # Replace each (…) Tj text run with /OCg-hidden drawing: simulate by
    # stripping text operators entirely (pypdf extracts the empty string).
    pattern = re.compile(rb"\([^\n()]*\) Tj T\*")
    return pattern.sub(b"", body)


def make_docx(sections: list[tuple[str | None, list[str]]]) -> bytes:
    """python-docx document: (heading, paragraph lines) pairs."""
    from docx import Document

    document = Document()
    for heading, paragraphs in sections:
        if heading:
            document.add_heading(heading, level=1)
        for line in paragraphs:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_table(heading: str, rows: list[list[str]]) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=2)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for col_index, cell in enumerate(row):
            table.rows[row_index].cells[col_index].text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip_bomb_docx() -> bytes:
    """A small zip pretending to hold word/document.xml but decompressing huge."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"<w:document xmlns:w='http://fake'/>")
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("bomb.bin", b"\x00" * (50 * 1024 * 1024))
    return buffer.getvalue()


def make_corrupt_pdf() -> bytes:
    return b"%PDF-1.4\nthis is not really a pdf stream" + b"garbage" * 200


def make_binary_junk() -> bytes:
    """Bytes that are neither PDF, nor a zip, nor text (a NUL byte proves binary)."""
    return b"RAW" + b"\x00" * 8 + b"BINARY" + b"\x01\x02\x03" * 40
